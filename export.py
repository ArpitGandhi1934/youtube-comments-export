#!/usr/bin/env python3
"""Export YouTube comments to CSV, Excel (.xlsx) or JSONL.

Calls the yugenox/youtube-scraper Actor on Apify for video URLs, a channel's latest videos or a
search, then writes one row per comment (and per reply, with --replies): author, text, likes,
reply count, date and the video it belongs to. The Actor bills per video, not per comment:
$0.002 per video whether it carries 10 comments or 5,000.

Run with --help for every option. Requires Python 3.10+ and `pip install -r requirements.txt`
(openpyxl is only needed for .xlsx). APIFY_TOKEN comes from the environment or a .env file next
to this script; --from needs no token.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

ACTOR = "yugenox/youtube-scraper"
PRICE_PER_VIDEO = 0.002  # USD per video row, every Apify plan (checked 2026-10-06)
START_FEE = 0.00005  # USD per run start (1 GB of memory)

# Column order for CSV, XLSX and JSONL.
COLUMNS = [
    "author", "text", "likes", "reply_count", "published", "published_approx", "is_reply",
    "parent_id", "author_is_creator", "is_pinned", "is_hearted", "author_is_verified", "comment_id",
    "video_id", "video_title", "video_url", "channel_name",
]
GUARDED = {"author", "text", "video_title", "channel_name"}

EPILOG = """\
price: $0.002 per video ($2.00 per 1,000 videos) on every Apify plan, with its comments and
replies included, plus $0.00005 per run (checked 2026-10-06).

examples:
  python export.py "https://www.youtube.com/watch?v=dQw4w9WgXcQ" --max-comments 1000 --out comments.xlsx
  python export.py --channel @NASA --videos 20 --max-comments 200 --replies 5
  python export.py --search "standing desk review" --videos 50 --out reviews.csv
  python export.py --file videos.txt --sort newest --max-comments 200 --append --out comments.csv
"""


def load_dotenv(path: Path = Path(__file__).with_name(".env")) -> None:
    """Minimal .env reader so the script has no extra dependency."""
    if not path.exists():
        return
    for line in path.read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def _field(obj, camel: str, snake: str):
    """apify-client v1/v2 return dicts, v3 returns model objects. Support both."""
    if obj is None:
        return None
    return obj.get(camel) if isinstance(obj, dict) else getattr(obj, snake, None)


def read_list(path: str) -> list[str]:
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return [s.strip() for s in lines if s.strip() and not s.strip().startswith("#")]


def is_single_video(s: str) -> bool:
    """A single video (counts as one) or a list (a playlist or channel URL, up to --videos each)."""
    return bool(re.fullmatch(r"[A-Za-z0-9_-]{11}", s) or re.search(r"youtu\.be/|[?&]v=|/shorts/|/embed/", s))


def handle(a):
    """The Actor's comment author is the @handle; the leading @ is dropped so that a CSV cell can't
    start with it (Excel can read a cell that starts with @ as a formula)."""
    return re.sub(r"^@", "", a) if isinstance(a, str) else None


def comment_row(c: dict, video: dict, parent_id) -> dict:
    approx = c.get("publishedAtApprox")
    return {
        "author": handle(c.get("author")),
        "text": c.get("text"),
        "likes": c.get("likes"),
        "reply_count": c.get("replyCount"),
        "published": c.get("publishedTime"),
        "published_approx": approx[:10] if isinstance(approx, str) else None,
        "is_reply": parent_id is not None,
        "parent_id": parent_id,
        "author_is_creator": c.get("authorIsCreator"),
        "is_pinned": c.get("isPinned"),
        "is_hearted": c.get("isHearted"),
        "author_is_verified": c.get("authorIsVerified"),
        "comment_id": c.get("commentId"),
        "video_id": video.get("id"),
        "video_title": video.get("title"),
        "video_url": video.get("url"),
        "channel_name": video.get("channelName"),
    }


def to_rows(r: dict) -> list[dict]:
    """One Actor row (a video with its comments) -> one output row per comment and per reply.
    A row this script wrote earlier passes through unchanged. Other rows (errors) give nothing."""
    if "comment_id" in r:
        return [{c: r.get(c) for c in COLUMNS}]
    if not isinstance(r.get("comments"), list):
        return []
    rows = []
    for c in r["comments"]:
        rows.append(comment_row(c, r, None))
        for reply in c.get("replies") or []:
            rows.append(comment_row(reply, r, c.get("commentId")))
    return rows


def csv_value(v, col: str):
    """Text that starts with = + - @, a tab or a carriage return gets a leading ' so Excel and Google
    Sheets show it as text instead of running it as a formula (comments are written by strangers)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if col in GUARDED and isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def existing_ids(path: str) -> set[str]:
    p = Path(path)
    if not p.exists() or p.stat().st_size == 0:
        return set()
    if path.lower().endswith(".jsonl"):
        with p.open(encoding="utf-8-sig") as f:
            return {json.loads(line).get("comment_id") for line in f if line.strip()} - {None}
    with p.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != COLUMNS:
            sys.exit(f"{path} has different columns than this script writes; use a new --out file")
        return {row["comment_id"] for row in reader if row.get("comment_id")}


def write_xlsx(path: str, rows: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("Writing .xlsx needs openpyxl: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Comments"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    text_cols = {"parent_id", "comment_id", "video_id"}
    for r in rows:
        ws.append([None] * len(COLUMNS))
        line = ws.max_row
        for i, col in enumerate(COLUMNS, start=1):
            v = r.get(col)
            if col == "published_approx" and v:
                # Midnight UTC of the approximate day; Excel has no time zones.
                v = datetime.fromisoformat(f"{v}T00:00:00")
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
            cell = ws.cell(row=line, column=i, value=v)
            if isinstance(v, str):
                cell.data_type = "s"  # never a formula, even if the comment starts with =
            if col == "published_approx" and v is not None:
                cell.number_format = "yyyy-mm-dd"
            elif col in text_cols:
                cell.number_format = "@"
            elif col == "text":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [22, 70, 9, 11, 14, 15, 9, 28, 16, 9, 10, 17, 28, 13, 40, 44, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def write_out(path: str, rows: list[dict], append: bool) -> None:
    low = path.lower()
    if low.endswith(".xlsx"):
        write_xlsx(path, rows)
        return
    p = Path(path)
    fresh = not append or not p.exists() or p.stat().st_size == 0
    if low.endswith(".jsonl"):
        with p.open("w" if fresh else "a", encoding="utf-8") as f:
            for r in rows:
                f.write(json.dumps(r, ensure_ascii=False, separators=(",", ":")) + "\n")
        return
    # UTF-8 with a BOM, so Excel shows emoji and accents correctly when it opens the file.
    with p.open("w" if fresh else "a", newline="", encoding="utf-8-sig" if fresh else "utf-8") as f:
        w = csv.writer(f)
        if fresh:
            w.writerow(COLUMNS)
        for r in rows:
            w.writerow([csv_value(r.get(c), c) for c in COLUMNS])


def read_from_file(path: str) -> list[dict]:
    text = Path(path).read_text(encoding="utf-8-sig").strip()
    if not text:
        return []
    if text.startswith("["):
        return json.loads(text)
    return [json.loads(line) for line in text.splitlines() if line.strip()]


def print_video_summary(items: list[dict]) -> None:
    """One line per video: how many comments came back, YouTube's total, and whether paging stopped early."""
    for v in items:
        if v.get("error"):
            src = v.get("fromYTUrl") or v.get("url") or v.get("input") or ""
            note = f" ({v['note']})" if v.get("note") else ""
            print(f"  {v['error']}  {src}{note}", file=sys.stderr)
            continue
        comments = v.get("comments")
        if not isinstance(comments, list):
            continue
        replies = sum(len(c.get("replies") or []) for c in comments)
        total = f" of {v['totalCommentsCount']:,} on YouTube" if v.get("totalCommentsCount") is not None else ""
        more = f" + {replies} replies" if replies else ""
        early = "  [stopped early]" if v.get("commentsTruncated") else ""
        print(f"  {len(comments)} comments{more}{total}  {v.get('title') or v.get('id')}{early}", file=sys.stderr)


def print_run_notes(client, run) -> None:
    message = _field(run, "statusMessage", "status_message")
    if message:
        print(f"  {message}", file=sys.stderr)
    # ERRORS (in the run's key-value store) lists inputs that returned nothing: an unknown
    # channel, an empty search, a private or deleted video. They are not charged.
    try:
        store_id = _field(run, "defaultKeyValueStoreId", "default_key_value_store_id")
        rec = client.key_value_store(store_id).get_record("ERRORS")
        value = rec.get("value") if isinstance(rec, dict) else getattr(rec, "value", None)
        for e in value if isinstance(value, list) else []:
            note = f" ({e['note']})" if e.get("note") else ""
            print(f"  {e.get('error', '?')}  {e.get('input') or e.get('url') or ''}{note}, not charged", file=sys.stderr)
    except Exception:  # the record is a convenience; the export doesn't depend on it
        pass


def main() -> None:
    p = argparse.ArgumentParser(
        prog="export.py",
        description="Export YouTube comments to CSV, Excel (.xlsx) or JSONL.",
        epilog=EPILOG,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("urls", nargs="*", metavar="URL",
                   help="video URLs (watch?v=, youtu.be/, /shorts/) or 11-character video ids; a playlist URL takes its first --videos videos, a channel URL its latest")
    v = p.add_argument_group("which videos")
    v.add_argument("--file", metavar="PATH", help="text file with one URL per line (# lines skipped)")
    v.add_argument("--channel", action="append", default=[], metavar="NAME", help="a channel's latest videos: name, @handle or channel URL; repeat for several")
    v.add_argument("--search", action="append", default=[], metavar="TERMS", help="videos that match a YouTube search; repeat for several")
    v.add_argument("--videos", type=int, default=10, metavar="N", help="videos per channel, playlist or search (default 10); above 40, a search keeps paging until it has N")
    v.add_argument("--since", metavar="WHEN", help='channel videos published on or after a date (2026-09-01) or within a period ("7 days"); older videos are skipped and not billed')
    c = p.add_argument_group("comments")
    c.add_argument("--max-comments", type=int, default=100, metavar="N", help="top-level comments per video, 1 to 5000 (default 100)")
    c.add_argument("--sort", choices=["top", "newest"], help='YouTube\'s "Top comments" order (default) or newest first')
    c.add_argument("--replies", type=int, default=0, metavar="N", help="also export up to N replies per comment (default 0); most-replied comments first, about 200 replies per video in total")
    o = p.add_argument_group("output")
    o.add_argument("--out", default="comments.csv", metavar="PATH", help="output path: .csv (default comments.csv), .xlsx or .jsonl")
    o.add_argument("--append", action="store_true", help="add only comments whose comment_id is not in --out yet (.csv or .jsonl)")
    o.add_argument("--max-charge", type=float, default=2.0, metavar="USD", help="hard cost cap for the run in USD, enforced by Apify (default 2, about 1,000 videos)")
    r = p.add_argument_group("re-export without a new run")
    r.add_argument("--dataset", metavar="ID", help="export the dataset of a run you already made (no new run starts)")
    r.add_argument("--from", dest="from_file", metavar="FILE", help="convert a JSON/JSONL file of the Actor's rows or of this script's rows (no token needed)")
    args = p.parse_args()

    urls = list(args.urls) + (read_list(args.file) if args.file else [])
    if not 1 <= args.videos <= 20000:
        p.error("--videos needs a whole number from 1 to 20000")
    if not 1 <= args.max_comments <= 5000:
        p.error("--max-comments needs a whole number from 1 to 5000")
    if not 0 <= args.replies <= 500:
        p.error("--replies needs a whole number from 0 to 500")
    if not args.max_charge > 0:
        p.error("--max-charge needs an amount in USD, for example 2")
    if not re.search(r"\.(csv|xlsx|jsonl)$", args.out, re.I):
        p.error("--out must end in .csv, .xlsx or .jsonl")
    if args.append and args.out.lower().endswith(".xlsx"):
        p.error("--append works with .csv or .jsonl (Excel opens the CSV)")
    fresh = bool(urls or args.channel or args.search)
    sources = sum(bool(x) for x in (fresh, args.dataset, args.from_file))
    if sources == 0:
        p.error("give video URLs, --file, --channel, --search, --dataset <id> or --from <file>")
    if sources > 1:
        p.error("use one of: URLs/--file/--channel/--search, --dataset, --from")

    if args.from_file:
        items = read_from_file(args.from_file)
    else:
        load_dotenv()
        token = os.environ.get("APIFY_TOKEN")
        if not token:
            sys.exit("Set APIFY_TOKEN (see .env.example), or use --from <file>")
        from apify_client import ApifyClient

        client = ApifyClient(token)
        dataset_id = args.dataset
        if not dataset_id:
            run_input: dict = {"maxItems": args.videos, "maxComments": args.max_comments}
            if urls:
                run_input["startUrls"] = urls
            if args.channel:
                run_input["channels"] = args.channel
            if args.search:
                run_input["searchTerms"] = args.search
            # A plain search stops at YouTube's first two result pages (about 40 videos).
            if args.search and args.videos > 40:
                run_input["deepSearch"] = True
            if args.since:
                run_input["publishedAfter"] = args.since
            if args.sort:
                run_input["commentsSortBy"] = args.sort
            if args.replies > 0:
                run_input["maxReplies"] = args.replies

            singles = sum(1 for u in urls if is_single_video(u))
            est = singles + (len(urls) - singles + len(args.channel) + len(args.search)) * args.videos
            fits = math.floor((args.max_charge - START_FEE) / PRICE_PER_VIDEO)
            if est > fits:
                print(f"Note: this can return up to {est} videos (${est * PRICE_PER_VIDEO:.2f}); "
                      f"--max-charge {args.max_charge:g} stops the run at about {fits}. Raise --max-charge to get them all.",
                      file=sys.stderr)
            what = ", ".join(x for x in (
                f"{len(urls)} URL(s)" if urls else "",
                f"{len(args.channel)} channel(s)" if args.channel else "",
                f"{len(args.search)} search(es)" if args.search else "",
            ) if x)
            print(f"Running {ACTOR} on {what}, up to {args.max_comments} comments per video ...", file=sys.stderr)
            # max_total_charge_usd caps what this run can cost you (pay-per-event Actors only).
            run = client.actor(ACTOR).call(
                run_input=run_input, max_total_charge_usd=Decimal(str(args.max_charge))
            )
            if run is None:
                sys.exit("The run did not finish")
            status = _field(run, "status", "status")
            print(f"Run {_field(run, 'id', 'id')} finished: {getattr(status, 'value', status)}", file=sys.stderr)
            print_run_notes(client, run)
            dataset_id = _field(run, "defaultDatasetId", "default_dataset_id")
        items = list(client.dataset(dataset_id).iterate_items())

    print_video_summary(items)
    rows = [row for it in items for row in to_rows(it)]
    skipped = 0
    if args.append:
        seen = existing_ids(args.out)
        before = len(rows)
        rows = [r for r in rows if not r.get("comment_id") or r["comment_id"] not in seen]
        skipped = before - len(rows)
    write_out(args.out, rows, args.append)

    verb = "Appended" if args.append else "Wrote"
    extra = f" ({skipped} already there)" if skipped else ""
    print(f"{verb} {len(rows)} comments to {args.out}{extra}", file=sys.stderr)


if __name__ == "__main__":
    main()
